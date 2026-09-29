from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "Template_TCC_Projeto_Integrado_NEES.pptx"
OUTPUT = BASE / "Apresentacao_TCC_IA_Bem_Estar_Estudantil.pptx"
NOTES_OUTPUT = BASE / "Roteiro_Apresentacao_TCC_20min.md"
NOTES = []

W = Inches(13.333)
H = Inches(7.5)
PURPLE = RGBColor(61, 32, 91)
VIOLET = RGBColor(104, 65, 151)
CYAN = RGBColor(21, 173, 186)
TEAL = RGBColor(16, 126, 137)
ORANGE = RGBColor(232, 132, 67)
RED = RGBColor(190, 65, 70)
INK = RGBColor(42, 42, 52)
MUTED = RGBColor(95, 95, 110)
LIGHT = RGBColor(246, 244, 249)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(224, 221, 230)


def remove_all_slides(prs):
    for slide_id in list(prs.slides._sldIdLst):
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill, radius=False, line=None, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line if line else fill
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Aptos", margin=0.04, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=18, color=INK, spacing=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.03)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.text = "•  " + p.text
    return box


def title(slide, number, heading, subtitle=None):
    rect(slide, 0, 0, 13.333, 0.16, PURPLE)
    rect(slide, 0, 7.30, 13.333, 0.20, CYAN)
    textbox(slide, f"{number:02d}", 0.55, 0.49, 0.55, 0.36, 13, CYAN, True)
    textbox(slide, heading, 1.10, 0.40, 11.55, 0.55, 25, PURPLE, True)
    if subtitle:
        textbox(slide, subtitle, 1.10, 0.97, 11.3, 0.34, 12, MUTED)


def card(slide, x, y, w, h, heading, body, accent=CYAN, body_size=16):
    rect(slide, x, y, w, h, LIGHT, True, GRAY)
    rect(slide, x, y, 0.08, h, accent, True, accent)
    textbox(slide, heading, x + 0.23, y + 0.18, w - 0.4, 0.35, 17, accent, True)
    textbox(slide, body, x + 0.23, y + 0.66, w - 0.42, h - 0.82, body_size, INK)


def arrow(slide, x1, y1, x2, y2, color=VIOLET, width=2.5):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    ln.line.end_arrowhead = True
    return ln


def add_notes(slide, text):
    # O template não fornece placeholder de notas para slides novos.
    # Mantemos o roteiro em arquivo próprio, mais prático para impressão/ensaio.
    NOTES.append(text)


def new_slide(prs, number, heading, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_bg(s)
    title(s, number, heading, subtitle)
    return s


def main():
    prs = Presentation(TEMPLATE)
    remove_all_slides(prs)
    prs.slide_width, prs.slide_height = W, H

    # 1 — Capa
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_bg(s, PURPLE)
    rect(s, 0, 0, 13.333, 7.5, PURPLE)
    rect(s, 0, 0, 0.20, 7.5, CYAN)
    rect(s, 8.8, -0.4, 5.2, 8.3, VIOLET, True, VIOLET, 18)
    textbox(s, "INTELIGÊNCIA ARTIFICIAL\nDE BAIXO CONSUMO", 0.85, 0.85, 7.7, 1.55, 28, WHITE, True)
    textbox(s, "para o monitoramento preventivo do bem-estar estudantil em escolas públicas de baixa conectividade",
            0.88, 2.54, 7.5, 1.15, 20, RGBColor(224, 240, 242))
    rect(s, 0.88, 4.12, 1.25, 0.09, CYAN)
    textbox(s, "Programa Escola Presente", 0.88, 4.38, 6.8, 0.42, 18, WHITE, True)
    textbox(s, "Ivo Calado · Rafael Durelli · Diego Dias", 0.88, 5.17, 7.1, 0.35, 15, WHITE)
    textbox(s, "Especialização em Gestão de Políticas Públicas Educacionais e Transformação Digital · UFAL · 2026",
            0.88, 5.63, 7.3, 0.70, 12, RGBColor(223, 216, 233))
    textbox(s, "TCC\nPROJETO\nINTEGRADO", 9.55, 2.05, 2.5, 1.7, 21, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_notes(s, "Tempo: 40 segundos. Apresente a equipe e a ideia central: usar IA de baixo consumo como apoio ao cuidado humano em escolas onde a conectividade e a oferta psicossocial são limitadas. Evite explicar a arquitetura ainda.")

    # 2 — problema
    s = new_slide(prs, 2, "O problema: sofrimento pouco visível e resposta desigual",
                  "A necessidade existe; a capacidade de identificar e acolher não está igualmente distribuída.")
    card(s, 0.75, 1.62, 3.75, 3.65, "1 em cada 7", "adolescentes de 10 a 19 anos vive com alguma condição de saúde mental no mundo.", PURPLE, 19)
    card(s, 4.78, 1.62, 3.75, 3.65, "50% / 40%", "metade dos jovens ouvidos pelo UNICEF sentiu necessidade de ajuda; 40% deles não procurou ninguém.", CYAN, 18)
    card(s, 8.81, 1.62, 3.75, 3.65, "96% ≠ acesso efetivo", "internet chegou às escolas, mas dispositivos e conexão para muitos acessos ainda são desiguais.", ORANGE, 18)
    rect(s, 1.55, 5.76, 10.25, 0.74, RGBColor(242, 237, 247), True, GRAY)
    textbox(s, "Dupla exclusão: menos apoio presencial + soluções digitais que dependem da nuvem.", 1.82, 5.94, 9.7, 0.35, 18, PURPLE, True, PP_ALIGN.CENTER)
    textbox(s, "Fontes: OMS (2025); UNICEF (2022); TIC Educação (2024).", 0.76, 6.73, 7.4, 0.24, 9, MUTED)
    add_notes(s, "Tempo: 2 minutos. Comece pelo direito à educação e permanência. Use os números apenas para dimensionar o problema. Destaque que baixa conectividade não causa sofrimento; ela reduz a capacidade de resposta e inviabiliza ferramentas dependentes da nuvem.")

    # 3 — pergunta e objetivos
    s = new_slide(prs, 3, "Pergunta orientadora e objetivo", "O desenho combina política pública, tecnologia apropriada e salvaguardas.")
    rect(s, 0.80, 1.55, 11.75, 1.58, PURPLE, True, PURPLE)
    textbox(s, "Como apoiar o monitoramento preventivo do bem-estar sem ampliar desigualdades, estigma, carga de trabalho ou riscos à privacidade?",
            1.18, 1.86, 11.0, 0.88, 22, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    card(s, 0.82, 3.62, 3.55, 2.10, "DIAGNOSTICAR", "necessidades, infraestrutura e capacidade real da rede", VIOLET, 16)
    card(s, 4.88, 3.62, 3.55, 2.10, "DESENHAR", "protocolo + aplicativo offline-first + revisão humana", CYAN, 16)
    card(s, 8.94, 3.62, 3.55, 2.10, "AVALIAR", "segurança, equidade, confiança, carga e resultados", ORANGE, 16)
    textbox(s, "Objetivo geral: estruturar uma política e um protocolo técnico-operacional replicável.", 1.05, 6.16, 11.2, 0.42, 17, PURPLE, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 1 minuto e 30 segundos. Leia a pergunta orientadora de forma pausada. Explique que o objetivo não é construir somente um aplicativo, mas uma política implementável com fluxo, governança e avaliação.")

    # 4 — solução
    s = new_slide(prs, 4, "A solução: Programa Escola Presente", "Um pacote técnico-educacional; a tecnologia é apenas uma de suas camadas.")
    card(s, 0.72, 1.52, 3.78, 4.72, "PROTOCOLO DE CUIDADO", "• escuta não julgadora\n• triagem e revisão humana\n• referência e contrarreferência\n• formação das equipes", PURPLE, 17)
    card(s, 4.78, 1.52, 3.78, 4.72, "APLICATIVO OFFLINE-FIRST", "• EDAE-A/DASS-21\n• processamento local\n• TinyML leve\n• alternativa em papel", CYAN, 17)
    card(s, 8.84, 1.52, 3.78, 4.72, "GOVERNANÇA E AVALIAÇÃO", "• dados mínimos e cifrados\n• painéis agregados\n• auditoria e contestação\n• critérios de suspensão", ORANGE, 17)
    textbox(s, "Não diagnostica · não prescreve · não substitui profissionais", 1.05, 6.55, 11.2, 0.34, 19, RED, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 2 minutos. Apresente as três camadas como inseparáveis. Explique que a EDAE-A é a adaptação brasileira da DASS-21 para adolescentes e será usada somente para rastreio. Reforce a frase final.")

    # 5 — fluxo operacional
    s = new_slide(prs, 5, "Como funciona: do convite ao cuidado", "O algoritmo recomenda prioridade; pessoas decidem e acolhem.")
    steps = [
        ("1", "ESCUTA", "21 itens + pedido direto de ajuda", PURPLE),
        ("2", "PROCESSAMENTO", "pontuação e TinyML no aparelho", VIOLET),
        ("3", "PRIORIDADE", "faixa auxiliar e fila protegida", CYAN),
        ("4", "REVISÃO HUMANA", "acolhimento e decisão profissional", TEAL),
        ("5", "REDE", "encaminhamento e retorno", ORANGE),
    ]
    for i, (n, h, b, c) in enumerate(steps):
        x = 0.62 + i * 2.52
        rect(s, x, 2.05, 2.08, 2.72, LIGHT, True, GRAY)
        rect(s, x + 0.72, 1.66, 0.64, 0.64, c, True, c)
        textbox(s, n, x + 0.72, 1.76, 0.64, 0.28, 16, WHITE, True, PP_ALIGN.CENTER)
        textbox(s, h, x + 0.18, 2.55, 1.72, 0.40, 14, c, True, PP_ALIGN.CENTER)
        textbox(s, b, x + 0.18, 3.12, 1.72, 1.12, 14, INK, False, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < 4:
            arrow(s, x + 2.10, 3.34, x + 2.43, 3.34, MUTED, 1.8)
    rect(s, 1.05, 5.38, 11.25, 0.88, RGBColor(238, 248, 249), True, CYAN)
    textbox(s, "Proteção imediata: diante de situação grave, aciona-se o protocolo humano — nunca se espera o algoritmo.", 1.34, 5.63, 10.65, 0.38, 17, TEAL, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 2 minutos. Percorra o fluxo da esquerda para a direita. Explique que qualquer estudante pode pedir conversa, independentemente do escore. Em situação grave, o protocolo humano é imediato. O registro deve conter apenas o necessário para acompanhar a providência.")

    # 6 — arquitetura e privacidade
    s = new_slide(prs, 6, "Arquitetura orientada à privacidade", "Processar localmente reduz dependência de conexão e exposição de dados.")
    rect(s, 0.78, 1.55, 4.00, 4.88, RGBColor(241, 238, 246), True, GRAY)
    textbox(s, "NO DISPOSITIVO", 1.15, 1.90, 3.25, 0.38, 18, PURPLE, True, PP_ALIGN.CENTER)
    for i, text in enumerate(["respostas individuais", "pontuação / inferência", "fila restrita"]):
        rect(s, 1.25, 2.55 + i * 0.92, 3.02, 0.62, WHITE, True, GRAY)
        textbox(s, text, 1.38, 2.72 + i * 0.92, 2.76, 0.26, 15, INK, i == 1, PP_ALIGN.CENTER)
    textbox(s, "cifrado · acesso por função · retenção mínima", 1.12, 5.55, 3.34, 0.52, 13, MUTED, False, PP_ALIGN.CENTER)
    arrow(s, 5.05, 3.77, 6.12, 3.77, CYAN, 3)
    textbox(s, "quando necessário", 4.85, 3.30, 1.48, 0.30, 11, MUTED, False, PP_ALIGN.CENTER)
    rect(s, 6.34, 1.55, 5.95, 4.88, RGBColor(238, 248, 249), True, GRAY)
    textbox(s, "PARA A GESTÃO", 6.74, 1.90, 5.15, 0.38, 18, TEAL, True, PP_ALIGN.CENTER)
    textbox(s, "INDICADORES\nAGREGADOS E\nANONIMIZADOS", 7.05, 2.68, 4.52, 1.55, 26, TEAL, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(s, "Sem nomes · sem respostas individuais\nsem uso punitivo, nota, seleção ou publicidade", 6.88, 4.76, 4.85, 0.76, 15, INK, False, PP_ALIGN.CENTER)
    textbox(s, "LGPD + proteção integral + supervisão humana", 2.25, 6.67, 8.9, 0.28, 16, PURPLE, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 1 minuto e 30 segundos. Contraste claramente os dois lados: dados individuais permanecem protegidos no dispositivo e só chegam à equipe autorizada; gestores veem tendências agregadas. Processamento local reduz riscos, mas não elimina a necessidade de controles e governança.")

    # 7 — implementação
    s = new_slide(prs, 7, "Implementação em ondas: testar antes de escalar", "Piloto proposto em aproximadamente 100 escolas do Norte e Nordeste.")
    phases = [
        ("PREPARAR", "Meses 1–5", "rede, RIPD, cocriação, protótipo", PURPLE),
        ("FORMAR", "Meses 5–6", "equipes e simulação dos fluxos", VIOLET),
        ("TESTAR", "Meses 6–8", "fluxo sem IA + modelo silencioso", CYAN),
        ("PILOTAR", "Meses 9–12", "uso assistido e expansão gradual", TEAL),
        ("AVALIAR", "Meses 13–18", "correções, relatório e decisão", ORANGE),
    ]
    y = 2.08
    for i, (h, period, body, c) in enumerate(phases):
        x = 0.60 + i * 2.53
        rect(s, x, y, 2.18, 2.65, LIGHT, True, GRAY)
        rect(s, x, y, 2.18, 0.58, c, True, c)
        textbox(s, h, x + 0.10, y + 0.13, 1.98, 0.28, 14, WHITE, True, PP_ALIGN.CENTER)
        textbox(s, period, x + 0.16, y + 0.83, 1.86, 0.28, 15, c, True, PP_ALIGN.CENTER)
        textbox(s, body, x + 0.18, y + 1.34, 1.82, 0.82, 14, INK, False, PP_ALIGN.CENTER)
        if i < 4:
            arrow(s, x + 2.18, y + 1.32, x + 2.47, y + 1.32, MUTED, 1.6)
    rect(s, 1.20, 5.38, 10.92, 0.85, RGBColor(250, 242, 235), True, ORANGE)
    textbox(s, "Expansão condicionada à capacidade de acolhimento, à segurança e à equidade — não ao número de instalações.", 1.50, 5.62, 10.32, 0.38, 16, ORANGE, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 2 minutos. Mostre o caráter adaptativo. Destaque o modo silencioso: o modelo gera saídas, mas elas não mudam a fila; são comparadas à avaliação humana. Só depois de segurança e equidade adequadas as faixas podem apoiar a priorização.")

    # 8 — atores
    s = new_slide(prs, 8, "Governança intersetorial", "A resposta depende da articulação entre educação, saúde, assistência social e comunidade.")
    rect(s, 4.62, 2.45, 4.10, 1.25, PURPLE, True, PURPLE)
    textbox(s, "ESTUDANTE E CUIDADO", 4.93, 2.86, 3.48, 0.40, 20, WHITE, True, PP_ALIGN.CENTER)
    actors = [
        (0.75, 1.55, "SECRETARIAS", "diretrizes, recursos e prestação de contas", VIOLET),
        (8.95, 1.55, "EQUIPE ESPECIALIZADA", "revisão, acolhimento e encaminhamento", CYAN),
        (0.75, 4.42, "ESCOLA", "rotina, proteção local e canal humano", TEAL),
        (8.95, 4.42, "FAMÍLIAS E CONSELHOS", "legitimidade, participação e controle social", ORANGE),
    ]
    for x, y0, h, b, c in actors:
        card(s, x, y0, 3.55, 1.52, h, b, c, 13)
    arrow(s, 4.30, 2.22, 4.78, 2.69, VIOLET, 2)
    arrow(s, 9.00, 2.22, 8.54, 2.69, CYAN, 2)
    arrow(s, 4.30, 5.02, 4.78, 3.48, TEAL, 2)
    arrow(s, 9.00, 5.02, 8.54, 3.48, ORANGE, 2)
    textbox(s, "Parceiros técnicos e universidades: desenvolver, formar, validar e auditar sob governança pública.", 1.55, 6.47, 10.25, 0.38, 15, PURPLE, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 1 minuto e 20 segundos. Enfatize que professores não serão terapeutas. A escola acolhe e aciona o fluxo; profissionais habilitados revisam e encaminham; secretarias garantem condições; estudantes e famílias participam do desenho e da avaliação.")

    # 9 — riscos
    s = new_slide(prs, 9, "Riscos principais e barreiras de segurança", "O piloto pode ser reduzido ou suspenso quando a proteção não estiver assegurada.")
    risks = [
        ("VAZAMENTO", "minimização, criptografia, acesso e plano de incidentes", PURPLE),
        ("ERROS DO MODELO", "canal humano, revisão, calibração e contestação", VIOLET),
        ("VIÉS", "auditoria por grupo e suspensão diante de disparidades", CYAN),
        ("REDE INSUFICIENTE", "mapear capacidade antes de oferecer triagem", TEAL),
        ("SOBRECARGA", "dimensionar demanda e reorganizar tarefas", ORANGE),
        ("USO PUNITIVO", "proibição normativa, segregação e auditoria", RED),
    ]
    for i, (h, b, c) in enumerate(risks):
        col, row = i % 3, i // 3
        card(s, 0.70 + col * 4.20, 1.55 + row * 2.33, 3.82, 1.88, h, b, c, 14)
    rect(s, 1.28, 6.27, 10.78, 0.58, RGBColor(250, 239, 239), True, RED)
    textbox(s, "Regra de ouro: nenhuma classificação pode negar acolhimento a quem pede ajuda.", 1.58, 6.42, 10.18, 0.30, 16, RED, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 2 minutos. Não leia todos os cartões. Agrupe-os em três tipos: proteção de dados, segurança clínica/algorítmica e capacidade institucional. Dê destaque à rede insuficiente: identificar sem conseguir responder pode gerar dano e frustração.")

    # 10 — avaliação
    s = new_slide(prs, 10, "Como saberemos se funciona?", "Avaliar o serviço de cuidado — não apenas a precisão do aplicativo.")
    headers = ["DIMENSÃO", "INDICADOR-CHAVE", "PERGUNTA"]
    widths = [2.25, 4.05, 5.15]
    xs = [0.72, 2.97, 7.02]
    for x, w, htxt in zip(xs, widths, headers):
        rect(s, x, 1.50, w, 0.54, PURPLE, False, PURPLE)
        textbox(s, htxt, x + 0.12, 1.65, w - 0.24, 0.26, 13, WHITE, True)
    rows = [
        ("Resposta", "tempo até a primeira escuta", "a rede acolhe em prazo útil?"),
        ("Bem-estar", "mudança nos escores", "há melhora consistente no tempo?"),
        ("Equidade", "diferenças de acesso e erro", "algum grupo é prejudicado?"),
        ("Proteção", "incidentes e acessos indevidos", "os dados permanecem seguros?"),
        ("Experiência", "confiança, estigma e utilidade", "a comunidade legitima o processo?"),
        ("Trabalho", "tempo adicional e sobrecarga", "a rotina é sustentável?"),
    ]
    for i, row in enumerate(rows):
        y = 2.07 + i * 0.65
        bg = WHITE if i % 2 == 0 else LIGHT
        for x, w, val in zip(xs, widths, row):
            rect(s, x, y, w, 0.63, bg, False, GRAY)
            textbox(s, val, x + 0.12, y + 0.16, w - 0.24, 0.26, 13, INK, i == 0 and x == xs[0])
    textbox(s, "Meta de 15% de redução dos escores = hipótese de avaliação, não promessa de impacto.", 1.10, 6.31, 11.1, 0.38, 17, ORANGE, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 2 minutos. Mostre que uma métrica isolada não decide a expansão. A meta de 15% vem da nota técnica e serve como hipótese de planejamento. A interpretação deve considerar perdas, composição da amostra, regressão à média e mudanças de contexto.")

    # 11 — contribuição e limites
    s = new_slide(prs, 11, "Contribuição, limites e próximos passos", "Transformação digital com valor público exige capacidade institucional.")
    card(s, 0.72, 1.52, 3.80, 4.75, "CONTRIBUIÇÃO", "Transforma uma recomendação técnica em política implementável: produto, fluxo, atores, indicadores, riscos e cronograma.", PURPLE, 17)
    card(s, 4.77, 1.52, 3.80, 4.75, "LIMITES", "Sem validação do piloto não há eficácia demonstrada. IA não corrige escassez de profissionais e pode reproduzir vieses.", ORANGE, 17)
    card(s, 8.82, 1.52, 3.80, 4.75, "PRÓXIMOS PASSOS", "Pactuar a rede, concluir o RIPD, cocriar o protótipo, testar acessibilidade e operar primeiro em modo silencioso.", CYAN, 17)
    textbox(s, "Tecnologia habilitadora, não solução autônoma.", 1.32, 6.58, 10.72, 0.34, 19, PURPLE, True, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 1 minuto e 30 segundos. Seja explícito sobre o caráter propositivo do TCC. A principal entrega é o desenho integrado. O protótipo e o piloto ainda precisam passar por validação ética, técnica e operacional.")

    # 12 — conclusão
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_bg(s, PURPLE)
    rect(s, 0, 0, 13.333, 7.5, PURPLE)
    rect(s, 0, 0, 0.20, 7.5, CYAN)
    textbox(s, "CONCLUSÃO", 0.90, 0.82, 3.4, 0.45, 16, CYAN, True)
    textbox(s, "Cuidado humano\norientado por evidências,\nmesmo quando a internet falha.", 0.88, 1.48, 8.35, 2.35, 30, WHITE, True)
    rect(s, 0.90, 4.21, 1.35, 0.09, CYAN)
    textbox(s, "A IA só gera valor público quando permanece subordinada à proteção integral, à decisão humana e à capacidade real da rede.", 0.90, 4.52, 8.45, 1.15, 19, RGBColor(228, 222, 237))
    rect(s, 9.70, 1.28, 2.15, 2.15, CYAN, True, CYAN)
    textbox(s, "?", 9.70, 1.56, 2.15, 1.05, 45, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(s, "Perguntas", 9.10, 4.38, 3.35, 0.60, 24, WHITE, True, PP_ALIGN.CENTER)
    textbox(s, "Obrigado.", 9.10, 5.12, 3.35, 0.42, 16, RGBColor(218, 210, 229), False, PP_ALIGN.CENTER)
    add_notes(s, "Tempo: 1 minuto. Retome a tese em uma frase: o valor da IA está em tornar a escuta e a priorização mais viáveis onde há poucos recursos, sem automatizar decisões de cuidado. Agradeça e abra para perguntas.")

    # Metadados e revisão final
    prs.core_properties.title = "Inteligência Artificial de Baixo Consumo para o Monitoramento Preventivo do Bem-Estar Estudantil"
    prs.core_properties.subject = "Apresentação de TCC — UFAL"
    prs.core_properties.author = "Ivo Calado; Rafael Durelli; Diego Dias"
    prs.core_properties.keywords = "UFAL, bem-estar estudantil, TinyML, offline-first, DASS-21, política pública"
    prs.save(OUTPUT)
    headings = [
        "Capa", "O problema", "Pergunta e objetivo", "A solução",
        "Fluxo de cuidado", "Arquitetura e privacidade", "Implementação",
        "Governança", "Riscos", "Avaliação", "Contribuição e limites", "Conclusão"
    ]
    roteiro = ["# Roteiro da apresentação — 20 minutos", ""]
    for i, (heading, note) in enumerate(zip(headings, NOTES), 1):
        roteiro.extend([f"## Slide {i} — {heading}", "", note, ""])
    roteiro.extend([
        "## Perguntas prováveis da banca", "",
        "- **Por que usar IA se a escala já tem pontuação?** A IA é proposta como apoio à calibração e priorização contextual, mas deve provar ganho sobre regras simples. Se não houver ganho seguro e equitativo, mantém-se a pontuação convencional e o fluxo humano.", "",
        "- **A aplicação faz diagnóstico?** Não. A EDAE-A/DASS-21 é usada para rastreio e monitoramento; qualquer decisão depende de escuta e julgamento profissional.", "",
        "- **De onde vem a meta de 15%?** Da nota técnica que fundamenta o projeto. É hipótese de avaliação e parâmetro de planejamento, não resultado obtido nem promessa de causalidade.", "",
        "- **Como agir sem psicólogo em todas as escolas?** O piloto só começa onde houver fluxo pactuado e capacidade mínima de acolhimento e encaminhamento. A tecnologia não substitui a ampliação de profissionais.", "",
        "- **Por que offline-first?** Para funcionar com conexão instável, reduzir transferência de dados sensíveis e evitar excluir justamente as escolas com menor infraestrutura.", ""
    ])
    NOTES_OUTPUT.write_text("\n".join(roteiro), encoding="utf-8")
    print(OUTPUT)
    print(NOTES_OUTPUT)


if __name__ == "__main__":
    main()
